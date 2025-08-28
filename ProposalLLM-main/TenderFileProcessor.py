import openpyxl
from docx import Document
import re
import os
import textract
import olefile
import json
from docx.text.paragraph import Paragraph
from docx.table import Table
from docx.oxml.text.paragraph import CT_P
from docx.oxml.table import CT_Tbl
from docx.document import Document

class TenderFileProcessor:
    """招标文件处理器"""
    
    def __init__(self):
        self.requirements = []
        self.technical_requirements = []
        self.business_requirements = []
        self.functional_requirements = []
        
    def extract_from_word(self, docx_path):
        """从Word文档中提取招标需求"""
        print(f"正在解析招标文件: {docx_path}")
        
        if docx_path.endswith('.doc'):
            # 使用系统命令处理.doc文件
            import subprocess
            try:
                # 尝试使用antiword
                result = subprocess.run(['antiword', docx_path], capture_output=True, text=True, encoding='utf-8')
                if result.returncode == 0:
                    text = result.stdout
                    return self._extract_from_text(text)
                else:
                    print(f"antiword解析失败: {result.stderr}")
                    # 尝试使用textract作为备选
                    try:
                        text = textract.process(docx_path).decode('utf-8')
                        return self._extract_from_text(text)
                    except Exception as e:
                        print(f"textract也失败了: {e}")
                        return []
            except Exception as e:
                print(f"处理.doc文件时出错: {e}")
                return []
        else:
            # 使用python-docx处理.docx文件
            doc = Document(docx_path)
            current_section = ""
            requirements = []
            
            for block in self._iter_block_items(doc):
                if isinstance(block, Paragraph):
                    text = block.text.strip()
                    if not text:
                        continue
                        
                    # 识别章节标题
                    if self._is_heading(block):
                        current_section = text
                        print(f"发现章节: {current_section}")
                        continue
                    
                    # 提取需求信息
                    if self._contains_requirement_keywords(text):
                        requirement = self._extract_requirement(text, current_section)
                        if requirement:
                            requirements.append(requirement)
                            print(f"提取需求: {requirement['content'][:50]}...")
            
            return requirements
    
    def extract_from_excel(self, excel_path):
        """从Excel文件中提取招标需求"""
        print(f"正在解析Excel文件: {excel_path}")
        
        wb = openpyxl.load_workbook(excel_path)
        sheet = wb.active
        requirements = []
        
        for row in sheet.iter_rows(min_row=2):
            # 读取A、B、C列作为需求信息
            a_content = row[0].value if row[0].value else ""
            b_content = row[1].value if row[1].value else ""
            c_content = row[2].value if row[2].value else ""
            
            if a_content or b_content or c_content:
                requirement = {
                    'category': 'technical',
                    'section': 'requirements',
                    'content': f"{a_content} {b_content} {c_content}".strip(),
                    'priority': 'high' if '必须' in a_content or '要求' in a_content else 'normal'
                }
                requirements.append(requirement)
        
        return requirements
    
    def extract_from_json(self, json_path):
        """从JSON文件中提取招标需求"""
        print(f"正在解析JSON招标文件: {json_path}")
        
        try:
            with open(json_path, 'r', encoding='utf-8') as f:
                data = json.load(f)
            
            requirements = []
            
            # 递归提取JSON中的所有需求信息
            self._extract_requirements_from_json(data, "", requirements)
            
            return requirements
            
        except Exception as e:
            print(f"解析JSON文件时出错: {e}")
            return []
    
    def _extract_requirements_from_json(self, data, parent_key, requirements):
        """递归提取JSON中的需求信息"""
        if isinstance(data, dict):
            for key, value in data.items():
                current_key = f"{parent_key}.{key}" if parent_key else key
                
                # 特别处理技术需求部分
                if key == "技术需求" or "技术" in key or "需求" in key:
                    if isinstance(value, dict):
                        # 技术需求是字典，提取所有子项
                        for sub_key, sub_value in value.items():
                            if isinstance(sub_value, dict):
                                # 递归处理子字典
                                self._extract_requirements_from_json(sub_value, f"{current_key}.{sub_key}", requirements)
                            elif isinstance(sub_value, str) and len(sub_value.strip()) > 10:
                                # 直接提取字符串内容作为需求
                                requirement = self._extract_requirement(sub_value, f"{current_key}.{sub_key}")
                                if requirement:
                                    requirements.append(requirement)
                                    print(f"提取技术需求: {requirement['content'][:50]}...")
                    elif isinstance(value, str) and len(value.strip()) > 10:
                        # 直接提取字符串内容作为需求
                        requirement = self._extract_requirement(value, current_key)
                        if requirement:
                            requirements.append(requirement)
                            print(f"提取技术需求: {requirement['content'][:50]}...")
                elif isinstance(value, str):
                    # 如果值是字符串，检查是否包含需求关键词
                    if self._contains_requirement_keywords(value) and len(value.strip()) > 10:
                        requirement = self._extract_requirement(value, current_key)
                        if requirement:
                            requirements.append(requirement)
                            print(f"提取需求: {requirement['content'][:50]}...")
                elif isinstance(value, dict):
                    # 如果值是字典，递归处理
                    self._extract_requirements_from_json(value, current_key, requirements)
                elif isinstance(value, list):
                    # 如果值是列表，递归处理每个元素
                    for i, item in enumerate(value):
                        if isinstance(item, dict):
                            self._extract_requirements_from_json(item, f"{current_key}[{i}]", requirements)
                        elif isinstance(item, str) and self._contains_requirement_keywords(item) and len(item.strip()) > 10:
                            requirement = self._extract_requirement(item, current_key)
                            if requirement:
                                requirements.append(requirement)
                                print(f"提取需求: {requirement['content'][:50]}...")
        elif isinstance(data, list):
            for i, item in enumerate(data):
                if isinstance(item, dict):
                    self._extract_requirements_from_json(item, f"{parent_key}[{i}]", requirements)
                elif isinstance(item, str) and self._contains_requirement_keywords(item) and len(item.strip()) > 10:
                    requirement = self._extract_requirement(item, parent_key)
                    if requirement:
                        requirements.append(requirement)
                        print(f"提取需求: {requirement['content'][:50]}...")
    
    def _iter_block_items(self, parent):
        """迭代文档元素"""
        parent_elm = parent.element.body if isinstance(parent, Document) else parent._element
        for child in parent_elm.iterchildren():
            if isinstance(child, CT_P):
                yield Paragraph(child, parent)
            elif isinstance(child, CT_Tbl):
                yield Table(child, parent)
    
    def _is_heading(self, paragraph):
        """判断是否为标题"""
        return paragraph.style.name.startswith('Heading') or \
               re.search(r'^\d+\.', paragraph.text) or \
               re.search(r'^[一二三四五六七八九十]+、', paragraph.text)
    
    def _contains_requirement_keywords(self, text):
        """检查文本是否包含需求关键词"""
        keywords = [
            '要求', '需求', '功能', '性能', '技术', '系统', '支持', '具备',
            '实现', '提供', '确保', '满足', '符合', '达到', '配置', '部署',
            '架构', '平台', '数据库', '接口', '协议', '标准', '规范',
            '建设', '构建', '开展', '建立', '设计', '开发', '集成',
            '管理', '服务', '应用', '数据', '处理', '分析', '监测',
            '监管', '平台', '系统', '服务', '应用', '支撑', '支持'
        ]
        return any(keyword in text for keyword in keywords)
    
    def _extract_from_text(self, text):
        """从文本中提取需求信息"""
        lines = text.split('\n')
        current_section = ""
        requirements = []
        
        for line in lines:
            line = line.strip()
            if not line:
                continue
            
            # 识别章节标题
            if self._is_text_heading(line):
                current_section = line
                print(f"发现章节: {current_section}")
                continue
            
            # 提取需求信息
            if self._contains_requirement_keywords(line):
                requirement = self._extract_requirement(line, current_section)
                if requirement:
                    requirements.append(requirement)
                    print(f"提取需求: {requirement['content'][:50]}...")
        
        return requirements
    
    def _is_text_heading(self, text):
        """判断文本是否为标题"""
        return re.search(r'^\d+\.', text) or \
               re.search(r'^[一二三四五六七八九十]+、', text) or \
               re.search(r'^[A-Z][A-Z\s]+$', text) or \
               re.search(r'^第.+章', text) or \
               re.search(r'^[A-Z]\d+\.', text)
    
    def _extract_requirement(self, text, section):
        """提取需求信息"""
        # 清理文本
        text = re.sub(r'\s+', ' ', text.strip())
        
        # 分类需求
        category = 'technical'
        if any(word in text for word in ['业务', '流程', '管理', '操作']):
            category = 'business'
        elif any(word in text for word in ['界面', '用户', '交互', '显示']):
            category = 'functional'
        
        return {
            'category': category,
            'section': section,
            'content': text,
            'priority': 'high' if any(word in text for word in ['必须', '要求', '关键']) else 'normal'
        }
    
    def process_tender_file(self, file_path):
        """处理招标文件"""
        if file_path.endswith('.docx') or file_path.endswith('.doc'):
            return self.extract_from_word(file_path)
        elif file_path.endswith('.xlsx') or file_path.endswith('.xls'):
            return self.extract_from_excel(file_path)
        elif file_path.endswith('.json'):
            return self.extract_from_json(file_path)
        else:
            raise ValueError(f"不支持的文件格式: {file_path}")
    
    def generate_requirements_summary(self, requirements):
        """生成需求摘要"""
        summary = {
            'total_requirements': len(requirements),
            'technical_requirements': len([r for r in requirements if r['category'] == 'technical']),
            'business_requirements': len([r for r in requirements if r['category'] == 'business']),
            'functional_requirements': len([r for r in requirements if r['category'] == 'functional']),
            'high_priority': len([r for r in requirements if r['priority'] == 'high']),
            'sections': {}
        }
        
        # 按章节分组
        for req in requirements:
            section = req['section']
            if section not in summary['sections']:
                summary['sections'][section] = []
            summary['sections'][section].append(req)
        
        return summary

def main():
    """测试招标文件处理"""
    processor = TenderFileProcessor()
    
    # 检查是否存在招标文件
    tender_files = []
    for file in os.listdir('.'):
        if file.endswith(('.docx', '.doc', '.xlsx', '.xls', '.json')) and '招标' in file:
            tender_files.append(file)
    
    if not tender_files:
        print("未找到招标文件，请将招标文件放在当前目录下，文件名包含'招标'字样")
        return
    
    print(f"找到招标文件: {tender_files}")
    
    # 处理第一个招标文件
    tender_file = tender_files[0]
    requirements = processor.process_tender_file(tender_file)
    
    # 生成摘要
    summary = processor.generate_requirements_summary(requirements)
    
    print(f"\n需求提取完成:")
    print(f"总需求数: {summary['total_requirements']}")
    print(f"技术需求: {summary['technical_requirements']}")
    print(f"业务需求: {summary['business_requirements']}")
    print(f"功能需求: {summary['functional_requirements']}")
    print(f"高优先级: {summary['high_priority']}")
    
    # 保存提取的需求
    with open('extracted_requirements.txt', 'w', encoding='utf-8') as f:
        f.write("招标需求提取结果\n")
        f.write("=" * 50 + "\n\n")
        
        for i, req in enumerate(requirements, 1):
            f.write(f"{i}. [{req['category']}] {req['content']}\n")
            f.write(f"   章节: {req['section']}\n")
            f.write(f"   优先级: {req['priority']}\n\n")
    
    print(f"\n需求已保存到: extracted_requirements.txt")
    
    return requirements

if __name__ == "__main__":
    main() 